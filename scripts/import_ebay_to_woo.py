"""
Import ebay_organized_products + catalog Excel into live WooCommerce.

Usage:
  python scripts/import_ebay_to_woo.py --dry-run
  python scripts/import_ebay_to_woo.py
  python scripts/import_ebay_to_woo.py --sku=ML-02
"""

from __future__ import annotations

import argparse
import json
import mimetypes
import os
import re
import sys
import time
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[1]
ORGANIZED = ROOT / "ebay_organized_products"
CATALOG_XLSX = ORGANIZED / "eBay_Products_Catalog.xlsx"
REPORT_XLSX = ROOT / "eBay-all-active-listings-report-2026-10-19-11342645128 (1).xlsx"
ENV_PATH = ROOT / "frontend" / ".env"
LOG_PATH = ROOT / "scripts" / "woo_import_log.json"

USD_FALLBACK_RATE = 95.0  # INR per USD if FX endpoint fails


def load_env(path: Path) -> dict[str, str]:
    env = {}
    if not path.exists():
        raise SystemExit(f"Missing {path}")
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        env[k.strip()] = v.strip().strip('"').strip("'")
    return env


def parse_sizes(raw) -> list[str]:
    if raw is None:
        return []
    if isinstance(raw, list):
        text = ",".join(str(x) for x in raw)
    else:
        text = str(raw)
    parts = []
    for token in re.split(r"[,;/|]", text):
        token = token.strip()
        if not token:
            continue
        if "size type" in token.lower() or "=" in token and "size=" not in token.lower():
            if token.lower().startswith("size="):
                token = token.split("=", 1)[1].strip()
            else:
                continue
        token = re.sub(r"^size\s*=\s*", "", token, flags=re.I).strip()
        token = token.replace("XXL", "2XL").replace("XXXL", "3XL").replace("XXXXL", "4XL")
        if re.fullmatch(r"(XS|S|M|L|XL|2XL|3XL|4XL|5XL)", token, re.I):
            parts.append(token.upper() if token.upper() != "XS" else "XS")
    # preserve order, unique
    seen = set()
    out = []
    for p in parts:
        key = p.upper()
        if key not in seen:
            seen.add(key)
            out.append(p.upper() if p.upper() != "XS" else "XS")
    return out


def guess_color(title: str, sku: str) -> str:
    n = f"{title} {sku}".lower()
    overrides = {
        "WL-06": "Black",
        "WS-05": "Brown",
    }
    if sku in overrides:
        return overrides[sku]
    rules = [
        ("black", "Black"),
        ("brown", "Brown"),
        ("tan", "Brown"),
        ("beige", "Brown"),
        ("camel", "Brown"),
        ("olive", "Green"),
        ("green", "Green"),
        ("grey", "Grey"),
        ("gray", "Grey"),
        ("blue", "Blue"),
        ("white", "White"),
        ("red", "Red"),
    ]
    for needle, color in rules:
        if needle in n:
            return color
    return "Brown"


def guess_material(title: str, subcategory: str) -> str:
    t = title.lower()
    if "wool" in t and "suede" not in t:
        return "Wool"
    if "sheepskin" in t:
        return "Sheepskin"
    if "suede" in t:
        return "Suede"
    if "leather" in t:
        return "Genuine Leather"
    n = subcategory.lower()
    if "wool" in n and "suede" not in n:
        return "Wool"
    if "suede" in n:
        return "Suede"
    return "Genuine Leather"


def map_leaf(title: str, subcategory: str, sku: str) -> str:
    # Style from product title only — folder names like "Suede Jackets & Shirts"
    # would otherwise mis-file every suede jacket as a shirt.
    t = title.lower()
    if sku == "WS-04" or ("wool" in t and "suede" not in t):
        return "wool-jackets"
    if "shirt" in t:
        return "leather-shirt-jackets"
    if "double breasted" in t or "trench" in t:
        return "double-breasted-coats"
    if "sheepskin" in t:
        return "sheepskin-jackets"
    if "biker" in t or "cafe racer" in t or "motorcycle" in t:
        return "biker-jackets"
    if "bomber" in t:
        return "bomber-jackets"
    if "suede" in t:
        return "suede-jackets"
    if "oversized" in t:
        return "oversized-jackets"
    return "leather-jackets"


def leaf_slug_for_dept(dept: str, leaf: str) -> str:
    if dept == "women" and leaf in {
        "leather-jackets",
        "suede-jackets",
        "biker-jackets",
        "bomber-jackets",
        "wool-jackets",
        "oversized-jackets",
    }:
        return f"women-{leaf}"
    return leaf


def usd_to_inr(usd: float, rate: float) -> str:
    raw = float(usd) * rate
    snapped = int(round(raw / 100.0) * 100) - 1
    if snapped < 999:
        snapped = int(round(raw))
    return str(snapped)


def parent_sku_from_variation(sku: str) -> str | None:
    m = re.match(r"^([A-Za-z]+)_(\d+)", str(sku))
    if not m:
        return None
    return f"{m.group(1).upper()}-{m.group(2)}"


def size_from_variation(sku: str, variation: str | None) -> str | None:
    if variation:
        m = re.search(r"Size\s*=\s*([A-Za-z0-9]+)", str(variation), re.I)
        if m:
            sizes = parse_sizes(m.group(1))
            if sizes:
                return sizes[0]
    tail = str(sku).split("_")[-1]
    sizes = parse_sizes(tail.replace("2X", "2XL").replace("3X", "3XL").replace("4X", "4XL"))
    return sizes[0] if sizes else None


class Woo:
    def __init__(self, base: str, ck: str, cs: str):
        self.base = base.rstrip("/")
        self.auth = (ck, cs)
        self.session = requests.Session()
        self.session.auth = self.auth
        self.session.headers.update({"User-Agent": "ZaynLeathers-Importer/1.0"})

    def wc(self, method: str, path: str, **kwargs):
        url = f"{self.base}/wp-json/wc/v3/{path.lstrip('/')}"
        r = self.session.request(method, url, timeout=120, **kwargs)
        if r.status_code >= 400:
            raise RuntimeError(f"{method} {path} -> {r.status_code} {r.text[:600]}")
        return r.json() if r.text else None

    def wp(self, method: str, path: str, **kwargs):
        url = f"{self.base}/wp-json/{path.lstrip('/')}"
        r = self.session.request(method, url, timeout=120, **kwargs)
        return r

    def get_all(self, path: str, extra: dict | None = None):
        page = 1
        out = []
        while True:
            params = {"per_page": 100, "page": page}
            if extra:
                params.update(extra)
            batch = self.wc("GET", path, params=params)
            if not batch:
                break
            out.extend(batch)
            if len(batch) < 100:
                break
            page += 1
        return out


def load_catalog() -> list[dict]:
    wb = openpyxl.load_workbook(CATALOG_XLSX, data_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not values[1]:
            continue
        item = dict(zip(headers, values))
        rows.append(
            {
                "item_id": str(item.get("Item ID") or ""),
                "sku": str(item.get("SKU") or "").strip(),
                "department_label": item.get("Department") or "",
                "subcategory": item.get("Subcategory") or "",
                "title": item.get("Product Title") or "",
                "price_usd": float(item.get("Price (USD)") or 0),
                "sizes_raw": item.get("Available Sizes") or "",
                "folder": item.get("Folder Path") or "",
                "ebay_url": item.get("eBay Listing URL") or "",
                "primary": item.get("Primary Image File") or "01_main.jpg",
            }
        )
    return rows


def load_variation_stock() -> dict[str, dict[str, int]]:
    """parent_sku -> {size: qty}"""
    stock: dict[str, dict[str, int]] = {}
    if not REPORT_XLSX.exists():
        return stock
    wb = openpyxl.load_workbook(REPORT_XLSX, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value or "").replace("\ufeff", "") for c in ws[1]]
    for values in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, values))
        sku = str(rec.get("Custom label (SKU)") or "").strip()
        if not sku or "-" in sku and "_" not in sku:
            continue
        parent = parent_sku_from_variation(sku)
        size = size_from_variation(sku, rec.get("Variation details"))
        if not parent or not size:
            continue
        qty = rec.get("Available quantity")
        try:
            qty_i = int(float(qty))
        except (TypeError, ValueError):
            qty_i = 1
        stock.setdefault(parent, {})[size] = max(qty_i, 0)
    return stock


def resolve_folder(row: dict) -> Path:
    rel = str(row["folder"]).replace("/", os.sep)
    direct = ORGANIZED / rel
    if direct.exists():
        return direct
    sku = row["sku"]
    for p in ORGANIZED.rglob("*"):
        if p.is_dir() and p.name.startswith(f"[{sku}]"):
            return p
    raise FileNotFoundError(f"No folder for {sku}")


def list_images(folder: Path, info: dict) -> list[Path]:
    names = []
    if info.get("primary_image"):
        names.append(info["primary_image"])
    for name in info.get("images") or []:
        if name not in names:
            names.append(name)
    files = []
    for name in names:
        path = folder / name
        if path.exists():
            files.append(path)
    if not files:
        files = sorted(
            [
                p
                for p in folder.iterdir()
                if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}
            ],
            key=lambda p: (0 if "main" in p.name else 1, p.name),
        )
    return files


def build_description(row: dict, color: str, material: str, sizes: list[str], sku: str) -> tuple[str, str]:
    short = f"{row['title']} in {color.lower()} {material.lower()}."
    long = (
        f"<p>Crafted for everyday wear with a clean modern cut. "
        f"Listed from the new Zayn Leathers collection.</p>"
    )
    return short, long


def ensure_categories(woo: Woo) -> dict[str, int]:
    existing = woo.get_all("products/categories")
    by_slug = {c["slug"]: c for c in existing}

    def upsert(name: str, slug: str, parent: int = 0) -> int:
        if slug in by_slug:
            return int(by_slug[slug]["id"])
        created = woo.wc(
            "POST",
            "products/categories",
            json={"name": name, "slug": slug, "parent": parent},
        )
        by_slug[slug] = created
        print(f"  + category {slug} ({created['id']})")
        return int(created["id"])

    ids = {}
    ids["men"] = upsert("Men", "men")
    ids["women"] = upsert("Women", "women")
    ids["accessories"] = upsert("Accessories", "accessories")
    ids["kids"] = upsert("Kids", "kids")

    leaves = {
        "men": [
            ("Leather Jackets", "leather-jackets"),
            ("Suede Jackets", "suede-jackets"),
            ("Biker Jackets", "biker-jackets"),
            ("Bomber Jackets", "bomber-jackets"),
            ("Sheepskin Jackets", "sheepskin-jackets"),
            ("Leather Shirt Jackets", "leather-shirt-jackets"),
            ("Double Breasted Coats", "double-breasted-coats"),
        ],
        "women": [
            ("Leather Jackets", "women-leather-jackets"),
            ("Suede Jackets", "women-suede-jackets"),
            ("Biker Jackets", "women-biker-jackets"),
            ("Bomber Jackets", "women-bomber-jackets"),
            ("Wool Jackets", "women-wool-jackets"),
            ("Oversized Jackets", "women-oversized-jackets"),
        ],
    }
    for dept, items in leaves.items():
        for name, slug in items:
            ids[slug] = upsert(name, slug, parent=ids[dept])
    return ids


def ensure_attribute(woo: Woo, name: str, slug: str, terms: list[str]) -> int:
    attrs = woo.get_all("products/attributes")
    found = next((a for a in attrs if a.get("slug") in {slug, f"pa_{slug}"} or a.get("name", "").lower() == name.lower()), None)
    if found:
        attr_id = int(found["id"])
    else:
        created = woo.wc(
            "POST",
            "products/attributes",
            json={"name": name, "slug": slug, "type": "select", "has_archives": True, "order_by": "menu_order"},
        )
        attr_id = int(created["id"])
        print(f"  + attribute {name} ({attr_id})")
    existing_terms = woo.get_all(f"products/attributes/{attr_id}/terms")
    have = {t["name"].lower() for t in existing_terms}
    for term in terms:
        if term.lower() in have:
            continue
        woo.wc("POST", f"products/attributes/{attr_id}/terms", json={"name": term})
        print(f"  + term {name}:{term}")
        have.add(term.lower())
    return attr_id


def upload_media(woo: Woo, path: Path, alt: str) -> int | None:
    mime = mimetypes.guess_type(path.name)[0] or "image/jpeg"
    with path.open("rb") as fh:
        files = {"file": (path.name, fh, mime)}
        data = {"alt_text": alt, "title": alt}
        r = woo.wp("POST", "wp/v2/media", files=files, data=data)
    if r.status_code < 300:
        return int(r.json()["id"])
    return None


def find_product(woo: Woo, sku: str) -> dict | None:
    found = woo.wc("GET", "products", params={"sku": sku})
    if found:
        return found[0]
    return None


def sync_variations(woo: Woo, product_id: int, size_attr_id: int, sizes: list[str], price: str, stock_map: dict[str, int], parent_sku: str):
    existing = woo.get_all(f"products/{product_id}/variations")
    by_size = {}
    for v in existing:
        for attr in v.get("attributes") or []:
            if str(attr.get("name", "")).lower() in {"size", "pa_size"} or attr.get("id") == size_attr_id:
                by_size[str(attr.get("option", "")).upper()] = v
    for size in sizes:
        qty = stock_map.get(size, 1)
        payload = {
            "sku": f"{parent_sku}-{size}",
            "regular_price": price,
            "manage_stock": True,
            "stock_quantity": qty,
            "stock_status": "instock" if qty > 0 else "outofstock",
            "attributes": [{"id": size_attr_id, "option": size}],
        }
        current = by_size.get(size)
        if current:
            # keep existing sku if already set
            if current.get("sku"):
                payload.pop("sku", None)
            woo.wc("PUT", f"products/{product_id}/variations/{current['id']}", json=payload)
        else:
            woo.wc("POST", f"products/{product_id}/variations", json=payload)
        time.sleep(0.15)


def import_one(woo: Woo, row: dict, cat_ids: dict[str, int], size_id: int, color_id: int, material_id: int, rate: float, stock: dict, dry: bool, only_sku: str | None) -> dict:
    sku = row["sku"]
    if only_sku and sku != only_sku:
        return {"sku": sku, "skipped": True}
    dept = "women" if "women" in str(row["department_label"]).lower() else "men"
    leaf = map_leaf(row["title"], row["subcategory"], sku)
    leaf_slug = leaf_slug_for_dept(dept, leaf)
    color = guess_color(row["title"], sku)
    material = guess_material(row["title"], row["subcategory"])
    sizes = parse_sizes(row["sizes_raw"])
    if not sizes:
        sizes = ["S", "M", "L", "XL", "2XL", "3XL"]
    price = usd_to_inr(row["price_usd"], rate)
    folder = resolve_folder(row)
    info_path = folder / "product_info.json"
    info = json.loads(info_path.read_text(encoding="utf-8")) if info_path.exists() else {}
    images = list_images(folder, info)
    stock_map = stock.get(sku, {})
    short, long = build_description(row, color, material, sizes, sku)

    result = {
        "sku": sku,
        "title": row["title"],
        "department": dept,
        "leaf": leaf_slug,
        "color": color,
        "material": material,
        "sizes": sizes,
        "price_inr": price,
        "price_usd": row["price_usd"],
        "images": [p.name for p in images],
        "folder": str(folder.relative_to(ORGANIZED)),
        "stock": {s: stock_map.get(s, 1) for s in sizes},
    }

    if dry:
        result["action"] = "dry-run"
        return result

    image_payload = []
    uploaded = 0
    for idx, img in enumerate(images):
        media_id = upload_media(woo, img, f"{sku} {idx+1:02d}")
        if media_id:
            image_payload.append({"id": media_id, "position": idx, "alt": f"{row['title']} {idx+1}"})
            uploaded += 1
        else:
            # fall back to eBay source URL if local media upload is blocked
            urls = info.get("image_urls") or []
            if idx < len(urls):
                image_payload.append({"src": urls[idx], "position": idx, "name": img.name, "alt": row["title"]})
        time.sleep(0.1)
    result["images_uploaded"] = uploaded
    result["images_payload"] = len(image_payload)

    cat_payload = [{"id": cat_ids[dept]}]
    if leaf_slug in cat_ids:
        cat_payload.append({"id": cat_ids[leaf_slug]})
    elif leaf in cat_ids:
        cat_payload.append({"id": cat_ids[leaf]})

    payload = {
        "name": row["title"],
        "type": "variable",
        "sku": sku,
        "status": "publish",
        "catalog_visibility": "visible",
        "featured": True,
        "description": long,
        "short_description": short,
        "categories": cat_payload,
        "tags": [{"name": "new-arrival"}, {"name": "new-collection"}],
        "attributes": [
            {
                "id": size_id,
                "visible": True,
                "variation": True,
                "options": sizes,
            },
            {
                "id": color_id,
                "visible": True,
                "variation": False,
                "options": [color],
            },
            {
                "id": material_id,
                "visible": True,
                "variation": False,
                "options": [material],
            },
        ],
        "images": image_payload,
        "meta_data": [
            {"key": "_ebay_item_id", "value": row["item_id"]},
            {"key": "_ebay_url", "value": row["ebay_url"]},
            {"key": "_source_sku", "value": sku},
        ],
    }

    existing = find_product(woo, sku)
    if existing and getattr(import_one, "skip_existing", False) and existing.get("images"):
        return {**result, "action": "exists", "woo_id": existing.get("id"), "permalink": existing.get("permalink")}
    if existing:
        product = woo.wc("PUT", f"products/{existing['id']}", json=payload)
        action = "updated"
    else:
        product = woo.wc("POST", "products", json=payload)
        action = "created"
    sync_variations(woo, int(product["id"]), size_id, sizes, price, stock_map, sku)
    result["action"] = action
    result["woo_id"] = product["id"]
    result["permalink"] = product.get("permalink")
    return result


def fetch_rate(base: str) -> float:
    try:
        r = requests.get(f"{base.rstrip('/')}/wp-json/zayn/v1/exchange-rates", timeout=20)
        data = r.json()
        usd = float(data.get("rates", {}).get("USD") or 0)
        if usd > 0:
            return 1.0 / usd
    except Exception:
        pass
    return USD_FALLBACK_RATE


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--sku", default="")
    parser.add_argument("--skip-existing", action="store_true")
    args = parser.parse_args()

    env = load_env(ENV_PATH)
    base = env.get("VITE_WC_URL")
    ck = env.get("VITE_WC_CONSUMER_KEY")
    cs = env.get("VITE_WC_CONSUMER_SECRET")
    if not (base and ck and cs):
        raise SystemExit("Missing Woo env keys in frontend/.env")

    catalog = load_catalog()
    stock = load_variation_stock()
    rate = fetch_rate(base)
    print(f"Catalog {len(catalog)} | FX {rate:.2f} INR/USD | dry={args.dry_run}")

    woo = None if args.dry_run else Woo(base, ck, cs)
    import_one.skip_existing = bool(args.skip_existing)
    cat_ids = size_id = color_id = material_id = None
    if not args.dry_run:
        print("Ensuring taxonomy...")
        cat_ids = ensure_categories(woo)
        size_id = ensure_attribute(woo, "Size", "size", ["XS", "S", "M", "L", "XL", "2XL", "3XL", "4XL"])
        color_id = ensure_attribute(woo, "Color", "color", ["Black", "Brown", "Blue", "Green", "White", "Red", "Grey"])
        material_id = ensure_attribute(woo, "Material", "material", ["Genuine Leather", "Suede", "Sheepskin", "Wool"])

    results = []
    for row in catalog:
        print(f"\n== {row['sku']} ==")
        try:
            if args.dry_run:
                rec = import_one(None, row, {}, 0, 0, 0, rate, stock, True, args.sku or None)
            else:
                rec = import_one(woo, row, cat_ids, size_id, color_id, material_id, rate, stock, False, args.sku or None)
            results.append(rec)
            print(json.dumps({k: rec[k] for k in rec if k != "images" or True}, default=str)[:700])
        except Exception as exc:
            rec = {"sku": row["sku"], "action": "error", "error": str(exc)[:800]}
            results.append(rec)
            print("ERROR", rec)
        time.sleep(0.2)

    LOG_PATH.write_text(json.dumps(results, indent=2), encoding="utf-8")
    created = sum(1 for r in results if r.get("action") == "created")
    updated = sum(1 for r in results if r.get("action") == "updated")
    errors = [r for r in results if r.get("action") == "error"]
    print(f"\nDone. created={created} updated={updated} errors={len(errors)} log={LOG_PATH}")
    if errors:
        sys.exit(2)


if __name__ == "__main__":
    main()
