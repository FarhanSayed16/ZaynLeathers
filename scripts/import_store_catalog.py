"""
Upload zaynleather_store_products into WooCommerce so products, photos,
categories, colours, and sizes are edited from the WordPress admin.

  python scripts/import_store_catalog.py --dry-run
  python scripts/import_store_catalog.py --limit 1
  python scripts/import_store_catalog.py

Already-published products with the same title are left in place.
New products use SKU ZL-{eBay item id}. Photos are copied into the
WordPress media library from the scraped image URLs.
"""

from __future__ import annotations

import argparse
import json
import re
import time
from collections import Counter
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parents[1]
CATALOG_DIR = ROOT / "zaynleather_store_products"
CATALOG_XLSX = CATALOG_DIR / "ZaynLeather_Master_Catalog.xlsx"
ENV_PATH = ROOT / "frontend" / ".env"
LOG_PATH = ROOT / "scripts" / "store_import_log.jsonl"
USD_FALLBACK_RATE = 95.0

LEAF_NAMES = {
    "leather-jackets": "Leather Jackets",
    "suede-jackets": "Suede Jackets",
    "biker-jackets": "Biker Jackets",
    "bomber-jackets": "Bomber Jackets",
    "sheepskin-jackets": "Sheepskin Jackets",
    "leather-shirt-jackets": "Leather Shirt Jackets",
    "double-breasted-coats": "Double Breasted Coats",
    "trench-coats": "Trench Coats",
    "leather-blazers": "Leather Blazers",
    "hooded-jackets": "Hooded Jackets",
    "leather-coats": "Leather Coats",
    "oversized-jackets": "Oversized Jackets",
    "leather-pants": "Leather Pants",
    "leather-cushion-covers": "Cushion Covers",
}

APPAREL_SIZES = ["S", "M", "L", "XL", "2XL", "3XL"]
PLUS_SIZES = ["2XL", "3XL", "4XL", "5XL"]
COLOR_RULES = [
    ("navy", "Blue"),
    ("burgundy", "Red"),
    ("maroon", "Red"),
    ("black", "Black"),
    ("white", "White"),
    ("cream", "White"),
    ("ivory", "White"),
    ("olive", "Green"),
    ("khaki", "Green"),
    ("green", "Green"),
    ("grey", "Grey"),
    ("gray", "Grey"),
    ("charcoal", "Grey"),
    ("blue", "Blue"),
    ("red", "Red"),
    ("wine", "Red"),
    ("tan", "Brown"),
    ("beige", "Brown"),
    ("camel", "Brown"),
    ("cognac", "Brown"),
    ("chocolate", "Brown"),
    ("brown", "Brown"),
]


def load_env(path: Path) -> dict[str, str]:
    env = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip().strip('"').strip("'")
    return env


def norm_title(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


def parse_usd(raw) -> float:
    if isinstance(raw, (int, float)):
        return float(raw)
    match = re.search(r"(\d+(?:\.\d+)?)", str(raw or "").replace(",", ""))
    return float(match.group(1)) if match else 0.0


def usd_to_inr(usd: float, rate: float) -> str:
    snapped = int(round((float(usd) * rate) / 100.0) * 100) - 1
    if snapped < 999:
        snapped = int(round(float(usd) * rate))
    return str(max(snapped, 0))


def parse_sizes(raw) -> list[str]:
    if not raw:
        return []
    found = []
    for token in re.split(r"[,;/|]", str(raw)):
        token = token.strip().upper().replace("XXXL", "3XL").replace("XXL", "2XL")
        token = token.replace("XXXXL", "4XL")
        if token in {"XS", "S", "M", "L", "XL", "2XL", "3XL", "4XL", "5XL"}:
            if token not in found:
                found.append(token)
    return found


def department_for(folder: str, title: str) -> str:
    blob = f"{folder} {title}".lower()
    if "pillow" in blob or "cushion" in blob:
        return "accessories"
    if "women" in blob or "ladies" in blob:
        return "women"
    return "men"


def leaves_for(folder: str, title: str) -> list[str]:
    blob = f"{folder} {title}".lower()
    if "pillow" in blob or "cushion" in blob:
        return ["leather-cushion-covers"]
    if "pant" in blob or "trouser" in blob:
        return ["leather-pants"]

    leaves = []

    def add(slug: str):
        if slug not in leaves:
            leaves.append(slug)

    if "blazer" in blob:
        add("leather-blazers")
    if "trench" in blob:
        add("trench-coats")
    if "hood" in blob:
        add("hooded-jackets")
    if "shirt" in blob:
        add("leather-shirt-jackets")
    if "double breasted" in blob or "double-breasted" in blob:
        add("double-breasted-coats")
    if "sheepskin" in blob or "shearling" in blob:
        add("sheepskin-jackets")
    if "biker" in blob or "cafe racer" in blob or "motorcycle" in blob:
        add("biker-jackets")
    if "bomber" in blob:
        add("bomber-jackets")
    if "suede" in blob:
        add("suede-jackets")
    if "oversized" in blob:
        add("oversized-jackets")
    if " coat" in f" {blob}" and "jacket" not in blob:
        add("leather-coats")
    if not leaves:
        add("leather-jackets")
    return leaves


def woo_slug(dept: str, leaf: str) -> str:
    if dept == "women":
        return f"women-{leaf}"
    return leaf


def color_for(title: str) -> str:
    text = title.lower()
    for needle, color in COLOR_RULES:
        if needle in text:
            return color
    return "Brown"


def material_for(folder: str, title: str) -> str:
    blob = f"{folder} {title}".lower()
    if "croc" in blob:
        return "Croco"
    if "faux" in blob or "pu leather" in blob:
        return "PU Leather"
    if "lambskin" in blob or "sheepskin" in blob or "shearling" in blob:
        return "Lambskin"
    if "suede" in blob:
        return "Suede"
    return "Genuine Leather"


def sizes_for(folder: str, title: str, raw) -> list[str]:
    parsed = parse_sizes(raw)
    if parsed:
        return parsed
    blob = f"{folder} {title}".lower()
    if "pillow" in blob or "cushion" in blob:
        return []
    if "plus size" in blob:
        return PLUS_SIZES
    return APPAREL_SIZES


class Woo:
    def __init__(self, base: str, key: str, secret: str):
        self.base = base.rstrip("/")
        self.session = requests.Session()
        self.session.auth = (key, secret)
        self.session.headers["User-Agent"] = "ZaynLeathers-Catalog/1.0"

    def wc(self, method: str, path: str, **kwargs):
        url = f"{self.base}/wp-json/wc/v3/{path.lstrip('/')}"
        last = None
        for attempt in range(4):
            response = self.session.request(method, url, timeout=240, **kwargs)
            if response.status_code in {429, 500, 502, 503, 504} and attempt < 3:
                time.sleep(2 + attempt * 3)
                last = response
                continue
            if response.status_code >= 400:
                raise RuntimeError(f"{method} {path} -> {response.status_code} {response.text[:500]}")
            return response.json() if response.text else None
        raise RuntimeError(f"{method} {path} -> {last.status_code if last else 'failed'}")

    def get_all(self, path: str, extra: dict | None = None):
        page = 1
        rows = []
        while True:
            params = {"per_page": 100, "page": page}
            if extra:
                params.update(extra)
            batch = self.wc("GET", path, params=params) or []
            rows.extend(batch)
            if len(batch) < 100:
                break
            page += 1
        return rows


def load_rows() -> list[dict]:
    book = openpyxl.load_workbook(CATALOG_XLSX, data_only=True, read_only=True)
    sheet = book.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not values or not values[0]:
            continue
        item = dict(zip(headers, values))
        item_id = str(item.get("Item ID") or "").strip()
        title = str(item.get("Product Title") or "").strip()
        folder = str(item.get("Category") or "").strip()
        if not item_id or not title:
            continue
        rows.append(
            {
                "item_id": item_id,
                "title": title,
                "folder": folder,
                "price_usd": parse_usd(item.get("Price")),
                "sizes_raw": item.get("Available Sizes") or "",
                "folder_path": str(item.get("Folder Path") or ""),
                "ebay_url": str(item.get("eBay Listing URL") or ""),
            }
        )
    return rows


def index_folders() -> dict[str, Path]:
    index = {}
    for category in CATALOG_DIR.iterdir():
        if not category.is_dir():
            continue
        for child in category.iterdir():
            if not child.is_dir() or not child.name.startswith("["):
                continue
            end = child.name.find("]")
            if end > 1:
                index[child.name[1:end]] = child
    return index


FOLDER_INDEX = None


def load_images(row: dict) -> list[str]:
    global FOLDER_INDEX
    if FOLDER_INDEX is None:
        FOLDER_INDEX = index_folders()
    folder = FOLDER_INDEX.get(row["item_id"])
    if folder is None:
        return []
    info_path = folder / "product_info.json"
    if not info_path.exists():
        return []
    info = json.loads(info_path.read_text(encoding="utf-8"))
    urls = []
    for url in info.get("image_urls") or []:
        if url and url not in urls:
            urls.append(url)
    return urls


def fetch_rate(base: str) -> float:
    try:
        response = requests.get(f"{base.rstrip('/')}/wp-json/zayn/v1/exchange-rates", timeout=20)
        usd = float(response.json().get("rates", {}).get("USD") or 0)
        if usd > 0:
            return 1.0 / usd
    except Exception:
        pass
    return USD_FALLBACK_RATE


def load_done() -> set[str]:
    done = set()
    if not LOG_PATH.exists():
        return done
    for line in LOG_PATH.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            rec = json.loads(line)
        except json.JSONDecodeError:
            continue
        if rec.get("action") in {"created", "exists", "skipped"} and rec.get("item_id"):
            done.add(str(rec["item_id"]))
    return done


def append_log(record: dict):
    LOG_PATH.parent.mkdir(parents=True, exist_ok=True)
    with LOG_PATH.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")


def ensure_categories(woo: Woo, needed: set[tuple[str, str]]) -> dict[str, int]:
    existing = woo.get_all("products/categories", {"hide_empty": False})
    by_slug = {cat["slug"]: cat for cat in existing}

    def upsert(name: str, slug: str, parent: int = 0) -> int:
        if slug in by_slug:
            return int(by_slug[slug]["id"])
        created = woo.wc("POST", "products/categories", json={"name": name, "slug": slug, "parent": parent})
        by_slug[slug] = created
        print(f"  + category {slug}")
        return int(created["id"])

    ids = {
        "men": upsert("Men", "men"),
        "women": upsert("Women", "women"),
        "accessories": upsert("Accessories", "accessories"),
        "kids": upsert("Kids", "kids"),
    }
    for dept, leaf in sorted(needed):
        slug = woo_slug(dept, leaf)
        parent = ids.get(dept, 0)
        ids[slug] = upsert(LEAF_NAMES.get(leaf, leaf.replace("-", " ").title()), slug, parent)
    return ids


def ensure_attribute(woo: Woo, name: str, slug: str, terms: list[str]) -> int:
    attrs = woo.get_all("products/attributes")
    found = next(
        (
            attr
            for attr in attrs
            if attr.get("slug") in {slug, f"pa_{slug}"} or str(attr.get("name", "")).lower() == name.lower()
        ),
        None,
    )
    if found:
        attr_id = int(found["id"])
    else:
        created = woo.wc(
            "POST",
            "products/attributes",
            json={"name": name, "slug": slug, "type": "select", "order_by": "menu_order", "has_archives": True},
        )
        attr_id = int(created["id"])
    have = {term["name"].lower() for term in woo.get_all(f"products/attributes/{attr_id}/terms")}
    for term in terms:
        if term.lower() in have:
            continue
        woo.wc("POST", f"products/attributes/{attr_id}/terms", json={"name": term})
        have.add(term.lower())
        print(f"  + {name}: {term}")
    return attr_id


def existing_titles(woo: Woo) -> set[str]:
    titles = set()
    for product in woo.get_all("products", {"status": "publish"}):
        titles.add(norm_title(product.get("name")))
        sku = str(product.get("sku") or "")
        if sku.startswith("ZL-"):
            titles.add(sku)
    return titles


def build_record(row: dict, rate: float) -> dict:
    dept = department_for(row["folder"], row["title"])
    leaves = leaves_for(row["folder"], row["title"])
    sizes = sizes_for(row["folder"], row["title"], row["sizes_raw"])
    return {
        "item_id": row["item_id"],
        "sku": f"ZL-{row['item_id']}",
        "title": row["title"],
        "department": dept,
        "leaves": leaves,
        "slugs": [woo_slug(dept, leaf) for leaf in leaves],
        "color": color_for(row["title"]),
        "material": material_for(row["folder"], row["title"]),
        "sizes": sizes,
        "price_inr": usd_to_inr(row["price_usd"], rate) if row["price_usd"] else "",
        "price_usd": row["price_usd"],
        "ebay_url": row["ebay_url"],
        "folder": row["folder"],
    }


def import_one(woo: Woo, row: dict, rec: dict, images: list[str], cat_ids: dict[str, int], size_id: int, color_id: int, material_id: int) -> dict:
    categories = [{"id": cat_ids[rec["department"]]}]
    for slug in rec["slugs"]:
        if slug in cat_ids:
            categories.append({"id": cat_ids[slug]})

    attributes = [
        {"id": color_id, "visible": True, "variation": False, "options": [rec["color"]]},
        {"id": material_id, "visible": True, "variation": False, "options": [rec["material"]]},
    ]
    product_type = "simple"
    if rec["sizes"]:
        product_type = "variable"
        attributes.insert(0, {"id": size_id, "visible": True, "variation": True, "options": rec["sizes"]})

    payload = {
        "name": rec["title"],
        "type": product_type,
        "sku": rec["sku"],
        "status": "publish",
        "catalog_visibility": "visible",
        "featured": False,
        "description": (
            f"<p>{rec['title']}. "
            f"{rec['material']} in {rec['color'].lower()}. "
            f"Handcrafted by Zayn Leathers and dispatched after inspection.</p>"
        ),
        "short_description": f"{rec['title']}.",
        "categories": categories,
        "attributes": attributes,
        "images": [{"src": url, "alt": rec["title"]} for url in images],
        "meta_data": [
            {"key": "_ebay_item_id", "value": rec["item_id"]},
            {"key": "_ebay_url", "value": rec["ebay_url"]},
            {"key": "_source_category", "value": rec["folder"]},
        ],
    }
    if product_type == "simple":
        payload["regular_price"] = rec["price_inr"]
        payload["manage_stock"] = True
        payload["stock_quantity"] = 1

    product = woo.wc("POST", "products", json=payload)
    if product_type == "variable":
        woo.wc(
            "POST",
            f"products/{product['id']}/variations/batch",
            json={
                "create": [
                    {
                        "sku": f"{rec['sku']}-{size}",
                        "regular_price": rec["price_inr"],
                        "manage_stock": True,
                        "stock_quantity": 1,
                        "attributes": [{"id": size_id, "option": size}],
                    }
                    for size in rec["sizes"]
                ]
            },
        )
    return {
        "item_id": rec["item_id"],
        "sku": rec["sku"],
        "action": "created",
        "woo_id": product.get("id"),
        "images": len(product.get("images") or []),
        "department": rec["department"],
        "slugs": rec["slugs"],
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--limit", type=int, default=0)
    args = parser.parse_args()

    env = load_env(ENV_PATH)
    base = env.get("VITE_WC_URL")
    key = env.get("VITE_WC_CONSUMER_KEY")
    secret = env.get("VITE_WC_CONSUMER_SECRET")
    if not (base and key and secret):
        raise SystemExit("Missing Woo keys in frontend/.env")

    rows = load_rows()
    rate = fetch_rate(base)
    print(f"Catalog {len(rows)} | FX {rate:.2f} INR/USD | dry={args.dry_run}")

    counts = Counter()
    missing_price = 0
    for row in rows:
        rec = build_record(row, rate)
        counts[f"{rec['department']}:{','.join(rec['slugs'])}"] += 1
        if not rec["price_inr"]:
            missing_price += 1
    print(f"Missing price: {missing_price}")
    for label, count in counts.most_common():
        print(f"{count:4}  {label}")

    if args.dry_run:
        missing_images = sum(1 for row in rows if not load_images(row))
        print("missing image sets", missing_images)
        return

    woo = Woo(base, key, secret)
    done = load_done()
    known_titles = existing_titles(woo)
    print(f"Already logged {len(done)} | live titles {len(known_titles)}")

    planned = []
    for row in rows:
        rec = build_record(row, rate)
        if rec["item_id"] in done:
            continue
        if rec["sku"] in known_titles or norm_title(rec["title"]) in known_titles:
            append_log({"item_id": rec["item_id"], "sku": rec["sku"], "action": "exists"})
            continue
        if not rec["price_inr"]:
            append_log({"item_id": rec["item_id"], "sku": rec["sku"], "action": "skipped", "error": "no price"})
            continue
        if not load_images(row):
            append_log({"item_id": rec["item_id"], "sku": rec["sku"], "action": "skipped", "error": "no images"})
            continue
        planned.append((row, rec))
    if args.limit:
        planned = planned[: args.limit]
    print(f"To create {len(planned)}")

    needed = {(rec["department"], leaf) for _, rec in planned for leaf in rec["leaves"]}
    # also include leaves from the whole catalog so empty styles still exist
    for row in rows:
        rec = build_record(row, rate)
        for leaf in rec["leaves"]:
            needed.add((rec["department"], leaf))

    print("Ensuring categories and attributes...")
    cat_ids = ensure_categories(woo, needed)
    size_id = ensure_attribute(woo, "Size", "size", ["XS", "S", "M", "L", "XL", "2XL", "3XL", "4XL", "5XL"])
    color_id = ensure_attribute(
        woo, "Color", "color", ["Black", "Brown", "Blue", "Green", "White", "Red", "Grey"]
    )
    material_id = ensure_attribute(
        woo, "Material", "material", ["Genuine Leather", "Suede", "Lambskin", "PU Leather", "Croco"]
    )

    created = 0
    errors = 0
    for index, (row, rec) in enumerate(planned, start=1):
        images = load_images(row)
        if not images:
            append_log({"item_id": rec["item_id"], "sku": rec["sku"], "action": "error", "error": "no images"})
            errors += 1
            print(f"[{index}/{len(planned)}] {rec['sku']} NO IMAGES")
            continue
        try:
            result = import_one(woo, row, rec, images, cat_ids, size_id, color_id, material_id)
            append_log(result)
            known_titles.add(norm_title(rec["title"]))
            known_titles.add(rec["sku"])
            created += 1
            print(f"[{index}/{len(planned)}] {rec['sku']} images={result['images']} id={result['woo_id']}")
        except Exception as exc:
            message = str(exc)
            plain = message.replace("\\/", "/")
            bad = [url for url in images if url in plain]
            if bad and len(bad) < len(images):
                try:
                    result = import_one(
                        woo, row, rec, [url for url in images if url not in bad],
                        cat_ids, size_id, color_id, material_id,
                    )
                    result["dropped_images"] = len(bad)
                    append_log(result)
                    known_titles.add(norm_title(rec["title"]))
                    known_titles.add(rec["sku"])
                    created += 1
                    print(f"[{index}/{len(planned)}] {rec['sku']} images={result['images']} id={result['woo_id']} dropped={len(bad)}")
                    time.sleep(0.15)
                    continue
                except Exception as retry_exc:
                    message = str(retry_exc)
            errors += 1
            append_log({"item_id": rec["item_id"], "sku": rec["sku"], "action": "error", "error": message[:500]})
            print(f"[{index}/{len(planned)}] {rec['sku']} ERROR {message[:180]}")
        time.sleep(0.15)

    print(f"Done. created={created} errors={errors} log={LOG_PATH}")


if __name__ == "__main__":
    main()
